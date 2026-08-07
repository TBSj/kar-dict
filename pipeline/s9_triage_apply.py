# -*- coding: utf-8 -*-
"""S9b-apply: merge triage verdicts + skeptic corrections into confirmed queues.

Input : checkpoints/shards/s9_triage_results.json  (workflow output)
Output: review_queues/rus_missing_confirmed.jsonl
        review_queues/rus_variants_confirmed.jsonl
        review_queues/rus_triage_rejected.jsonl
        review_queues/rus_check_review.xlsx  (для лингвиста)
        appends stats to qa/s9_rus_check.md
"""
import json, sys, io
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
CK = ROOT / "checkpoints"
RQ = CK / "review_queues"

YELLOW = PatternFill("solid", start_color="FFF3B0")
GREEN = PatternFill("solid", start_color="D9EAD3")
HDR = Font(bold=True)


def main():
    results = json.load(open(CK / "shards" / "s9_triage_results.json", encoding="utf-8"))

    items = {}   # form -> merged record
    n_corr = 0
    for r in results:
        batch = {c["form"]: c for c in json.load(open(r["batch"], encoding="utf-8"))}
        for it in r["triage"]["items"]:
            rec = dict(it)
            src = batch.get(it["form"], {})
            rec["form_src"] = src.get("form_src", it["form"])
            rec["queue"] = src.get("queue", "?")
            rec["count_cite"] = src.get("count_cite", 0)
            rec["count_ctx"] = src.get("count_ctx", 0)
            rec["near"] = src.get("near", [])
            rec["contexts"] = src.get("contexts", [])
            rec["corrected"] = False
            items[it["form"]] = rec
        for c in r["verify"].get("corrections", []):
            rec = items.get(c["form"])
            if not rec:
                continue
            rec["verdict"] = c["new_verdict"]
            if c.get("new_canonical"):
                rec["canonical"] = c["new_canonical"]
            rec["corrected"] = True
            rec["correction_reason"] = c.get("reason", "")
            n_corr += 1

    missing = [r for r in items.values() if r["verdict"] == "missing_word"]
    variants = [r for r in items.values() if r["verdict"] == "variant_spelling"]
    rejected = [r for r in items.values() if r["verdict"] not in ("missing_word", "variant_spelling")]
    key = lambda r: (-(r["count_cite"] * 3 + r["count_ctx"]), r["form"])
    missing.sort(key=key); variants.sort(key=key); rejected.sort(key=key)

    def dump(path, rows):
        with open(RQ / path, "w", encoding="utf-8") as f:
            for r in rows:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")

    dump("rus_missing_confirmed.jsonl", missing)
    dump("rus_variants_confirmed.jsonl", variants)
    dump("rus_triage_rejected.jsonl", rejected)

    wb = Workbook()
    ws = wb.active
    ws.title = "Пропуски"
    ws.append(["слово (canonical)", "как в источнике", "перевод (из контекстов)",
               "cite", "ctx", "conf", "скептик правил", "контексты", "вердикт лингвиста"])
    for c in ws[1]:
        c.font = HDR
    for r in missing:
        ctx = " || ".join(f"[{c['rus']}] {c['snippet']}" for c in r["contexts"])[:900]
        ws.append([r.get("canonical") or r["form"], r["form_src"],
                   r.get("ru_gloss", ""), r["count_cite"], r["count_ctx"],
                   round(r.get("conf", 0), 2), "да" if r["corrected"] else "", ctx, ""])
        row = ws.max_row
        if r.get("conf", 0) < 0.7 or r["corrected"]:
            for c in ws[row][:7]:
                c.fill = YELLOW
        else:
            ws.cell(row, 1).fill = GREEN
    ws2 = wb.create_sheet("Варианты")
    ws2.append(["форма", "слово корпуса (base)", "перевод", "cite", "ctx", "conf",
                "скептик правил", "контексты", "вердикт лингвиста"])
    for c in ws2[1]:
        c.font = HDR
    for r in variants:
        ctx = " || ".join(f"[{c['rus']}] {c['snippet']}" for c in r["contexts"])[:900]
        ws2.append([r["form"], r.get("base", "") or ", ".join(r.get("near", []))
                    or r.get("correction_reason", "")[:90],
                    r.get("ru_gloss", ""), r["count_cite"], r["count_ctx"],
                    round(r.get("conf", 0), 2), "да" if r["corrected"] else "", ctx, ""])
        if r.get("conf", 0) < 0.7 or r["corrected"]:
            for c in ws2[ws2.max_row][:7]:
                c.fill = YELLOW
    ws3 = wb.create_sheet("Отклонено")
    ws3.append(["форма", "вердикт", "base", "причина скептика", "cite", "ctx"])
    for c in ws3[1]:
        c.font = HDR
    for r in rejected:
        ws3.append([r["form"], r["verdict"], r.get("base", ""),
                    r.get("correction_reason", ""), r["count_cite"], r["count_ctx"]])
    for w in (ws, ws2, ws3):
        for col, width in zip("ABCDEFGHI", (26, 22, 34, 6, 6, 6, 12, 90, 24)):
            w.column_dimensions[col].width = width
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=(c.column_letter in "CH"))
    wb.save(RQ / "rus_check_review.xlsx")

    qa_path = CK / "qa" / "s9_rus_check.md"
    txt = open(qa_path, encoding="utf-8").read()
    cut = txt.find("\n## Триаж агентами")
    if cut != -1:
        open(qa_path, "w", encoding="utf-8").write(txt[:cut])
    with open(qa_path, "a", encoding="utf-8") as f:
        f.write(f"\n## Триаж агентами (600 кандидатов, 6 батчей + 6 скептиков)\n\n")
        f.write(f"- подтверждено пропущенных слов: {len(missing)}\n")
        f.write(f"- подтверждено вариантов написания: {len(variants)}\n")
        f.write(f"- отклонено (словоформы/русизмы/обрывки/фразы): {len(rejected)}\n")
        f.write(f"- поправок скептиков: {n_corr}\n")
        f.write(f"- на просмотр лингвисту: review_queues/rus_check_review.xlsx\n")
    print(f"missing {len(missing)}, variants {len(variants)}, rejected {len(rejected)}, "
          f"corrections {n_corr} -> rus_check_review.xlsx")


if __name__ == "__main__":
    main()
