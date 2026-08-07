# -*- coding: utf-8 -*-
"""Единая книга ручной доработки: manual_review_2.xlsx.

Собирает ВСЕ открытые очереди в один файл:
  1. Пропуски RU-КБ (222, из s9-триажа)
  2. Варианты RU-КБ (65)
  3. Варианты корпуса (739, s8-детектор: вариант vs синоним)
  4. needs_review новые (появились после прошлой книги manual_review.xlsx)
  5. Пометы-хвост (732 нераспознанных токена -> фасет)
  6. дж-ж остаток (2 имени собственных)
  7. Фраземы ромб (выборка из 3 625 статей с «◊» + вопрос о политике)
  8. Прочее (справочно: потери ТСКБЯ, акъыл/эт-джен)
  9. Отклонено RU-КБ (313, справочно)

Конвенция прежняя: жёлтая колонка = вердикт, пусто = согласен.
"""
import json, sys, io
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
CK = ROOT / "checkpoints"
RQ = CK / "review_queues"

YELLOW = PatternFill("solid", start_color="FFF3B0")
GREEN = PatternFill("solid", start_color="D9EAD3")
GREY = PatternFill("solid", start_color="EFEFEF")
HDR = Font(bold=True)
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")


def jl(p):
    return [json.loads(l) for l in open(p, encoding="utf-8")]


def short(s, n=220):
    s = s or ""
    return s if len(s) <= n else s[:n] + "…"


def style(ws, widths, wrap_cols, verdict_col=None):
    for c in ws[1]:
        c.font = HDR
        c.fill = GREY
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP if c.column_letter in wrap_cols else TOP
            if verdict_col and c.column_letter == verdict_col:
                c.fill = YELLOW
    ws.freeze_panes = "A2"


def main():
    entries = jl(CK / "07_final.jsonl")
    by_word = {}
    for e in entries:
        by_word.setdefault(e["word"], []).append(e)

    wb = Workbook()

    # ---------- Инструкция ----------
    ws = wb.active
    ws.title = "Инструкция"
    lines = [
        ("Единая книга ручной доработки объединённого словаря (все открытые очереди)",),
        ("",),
        ("Общая конвенция: ЖЁЛТАЯ колонка = твой вердикт. Пусто = согласен с тем, что есть.",),
        ("Зелёная заливка = высокая уверенность автоматики; жёлтая заливка строки = сомнительно.",),
        ("",),
        ("1_Пропуски_ру-кб — 222 слова из русско-КБ словаря, которых нет в корпусе.",),
        ("   Вердикт: пусто = добавить как есть; «нет» = не добавлять; или впиши правку формы/перевода.",),
        ("2_Варианты_ру-кб — 65 форм из русско-КБ словаря, похожих на слова корпуса.",),
        ("   Вердикт: пусто = это вариант написания (свяжем); «отд» = отдельное слово; «нет» = мусор.",),
        ("3_Варианты_корпуса — 739 автоматических связей вариант↔слово внутри корпуса.",),
        ("   Автоматика не отличает вариант от синонима. Вердикт: пусто = вариант; «син» = синонимы,",),
        ("   связь убрать; «нет» = ложная связь.",),
        ("4_needs_review_новые — статьи, помеченные к проверке ПОСЛЕ прошлой книги (ab-раунд, кь/гь).",),
        ("   Вердикт: пусто = ок, снять флаг; или впиши, что исправить.",),
        ("5_Пометы_хвост — 732 нераспознанных токена помет. Впиши фасет: pos / dialect / style /",),
        ("   domain / gram_note / мусор / оставить. Можно группами (одинаковое — одно слово).",),
        ("6_дж-ж_остаток — 2 имени собственных: оставить «Дж» или перевести в j?",),
        ("7_Фраземы_ромб — ◊ в ТСКБЯ: где фразема, где цитата. 30 примеров из 3 625 статей.",),
        ("   Нужно ПРАВИЛО (напиши в A2 листа), не построчные вердикты.",),
        ("8_Прочее — справка о задачах, которые делаю я (потери ТСКБЯ, акъыл/эт-джен).",),
        ("9_Отклонено_ру-кб — что триаж отклонил (справочно; глянь по диагонали — вдруг зря).",),
        ("",),
        ("Готовый файл сохрани как manual_review_2_ab.xlsx и скинь мне.",),
    ]
    for row in lines:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 110

    # ---------- 1. Пропуски RU-КБ ----------
    ws = wb.create_sheet("1_Пропуски_ру-кб")
    ws.append(["слово", "как в источнике", "перевод (из контекстов)", "cite", "ctx",
               "conf", "скептик правил", "контексты употребления", "ВЕРДИКТ"])
    for r in jl(RQ / "rus_missing_confirmed.jsonl"):
        ctx = " || ".join(f"[{c['rus']}] {c['snippet']}" for c in r["contexts"])[:900]
        ws.append([r.get("canonical") or r["form"], r["form_src"], r.get("ru_gloss", ""),
                   r["count_cite"], r["count_ctx"], round(r.get("conf", 0), 2),
                   r.get("correction_reason", "да")[:90] if r["corrected"] else "", ctx, ""])
        if r.get("conf", 0) >= 0.8 and not r["corrected"]:
            ws.cell(ws.max_row, 1).fill = GREEN
        elif r.get("conf", 0) < 0.7:
            for c in ws[ws.max_row][:7]:
                c.fill = YELLOW
    style(ws, {"A": 26, "B": 22, "C": 36, "D": 6, "E": 6, "F": 6, "G": 24, "H": 80, "I": 22},
          "CH", "I")

    # ---------- 2. Варианты RU-КБ ----------
    ws = wb.create_sheet("2_Варианты_ру-кб")
    ws.append(["форма (ру-кб словарь)", "слово корпуса", "перевод", "cite", "ctx",
               "conf", "контексты", "ВЕРДИКТ"])
    for r in jl(RQ / "rus_variants_confirmed.jsonl"):
        ctx = " || ".join(f"[{c['rus']}] {c['snippet']}" for c in r["contexts"])[:700]
        base = r.get("base", "") or ", ".join(r.get("near", [])) or r.get("correction_reason", "")[:90]
        ws.append([r["form"], base, r.get("ru_gloss", ""), r["count_cite"], r["count_ctx"],
                   round(r.get("conf", 0), 2), ctx, ""])
        if r.get("conf", 0) < 0.7:
            for c in ws[ws.max_row][:6]:
                c.fill = YELLOW
    style(ws, {"A": 24, "B": 30, "C": 30, "D": 6, "E": 6, "F": 6, "G": 70, "H": 20}, "CG", "H")

    # ---------- 3. Варианты корпуса ----------
    ws = wb.create_sheet("3_Варианты_корпуса")
    ws.append(["вариант", "основное слово", "def варианта", "def основного", "ВЕРДИКТ"])
    for r in jl(RQ / "variants_detected.jsonl"):
        d_v = by_word.get(r["variant"], [{}])[0]
        d_o = by_word.get(r["of"], [{}])[0]
        ws.append([r["variant"], r["of"],
                   short(d_v.get("def1") or d_v.get("def2")),
                   short(d_o.get("def1") or d_o.get("def2")), ""])
    style(ws, {"A": 20, "B": 20, "C": 55, "D": 55, "E": 16}, "CD", "E")

    # ---------- 4. needs_review новые ----------
    old = load_workbook(RQ / "manual_review.xlsx", read_only=True)
    seen_words = set()
    for name in ("needs_review", "split_unclear"):
        for row in old[name].iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:
                seen_words.add(str(row[0]))
    old.close()
    ws = wb.create_sheet("4_needs_review_новые")
    ws.append(["слово", "словари", "пометы", "def1", "def2 (ТСКБЯ)", "misc (сокр.)",
               "флаги", "ВЕРДИКТ"])
    n_new = 0
    for e in entries:
        if "needs_review" not in e.get("flags", []) or e["word"] in seen_words:
            continue
        misc = short((e.get("misc1") or "") + (" ▏" + e.get("misc2", "") if e.get("misc2") else ""), 300)
        ws.append([e["word"], e.get("dict_id", ""), e.get("tags", ""),
                   short(e.get("def1")), short(e.get("def2")), misc,
                   ";".join(f for f in e.get("flags", []) if f != "needs_review"), ""])
        n_new += 1
    style(ws, {"A": 18, "B": 8, "C": 16, "D": 45, "E": 45, "F": 40, "G": 26, "H": 20}, "DEF", "H")

    # ---------- 5. Пометы-хвост ----------
    tags = json.load(open(CK / "qa" / "s8_unmapped_tags.json", encoding="utf-8"))
    examples = {}
    for e in entries:
        for t in e.get("tags_other") or []:
            ex = examples.setdefault(t, [])
            if len(ex) < 3:
                ex.append(e["word"])
    ws = wb.create_sheet("5_Пометы_хвост")
    ws.append(["токен", "частота", "примеры слов", "ФАСЕТ (pos/dialect/style/domain/gram_note/мусор/оставить)"])
    for tok, cnt in sorted(tags.items(), key=lambda kv: -kv[1]):
        ws.append([tok, cnt, ", ".join(examples.get(tok, [])), ""])
    style(ws, {"A": 24, "B": 9, "C": 40, "D": 45}, "C", "D")

    # ---------- 6. дж-ж остаток ----------
    ws = wb.create_sheet("6_дж-ж_остаток")
    ws.append(["слово", "статья", "вопрос", "ВЕРДИКТ"])
    for r in jl(RQ / "zh_dzh_unmatched.jsonl"):
        ws.append([r["word"], short(r.get("def", ""), 300),
                   "имя собственное — оставить «Дж» или перевести в j?", ""])
    style(ws, {"A": 16, "B": 60, "C": 40, "D": 20}, "BC", "D")

    # ---------- 7. Фраземы ◊ ----------
    ws = wb.create_sheet("7_Фраземы_ромб")
    ws.append(["ПРАВИЛО (впиши сюда): как отличать фразему от цитаты после ◊?"])
    ws.append([""])
    ws["A2"].fill = YELLOW
    ws.append(["слово", "фрагмент с ◊"])
    dia = [e for e in entries if "◊" in (e.get("def2", "") + e.get("misc2", ""))]
    step = max(1, len(dia) // 30)
    for e in dia[::step][:30]:
        txt = e.get("def2", "") + " ▏ " + e.get("misc2", "")
        i = txt.find("◊")
        ws.append([e["word"], short(txt[max(0, i - 60):i + 260], 320)])
    for c in ws[3]:
        c.font = HDR
    ws["A1"].font = HDR
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = WRAP if c.column_letter == "B" else TOP

    # ---------- 8. Прочее ----------
    ws = wb.create_sheet("8_Прочее")
    ws.append(["задача", "статус / что нужно от тебя"])
    ws.append(["~37 потерянных заголовков ТСКБЯ + 7 склеек",
               "делаю я: точечный ре-парсинг по PDF; от тебя — подтверждение результата, когда покажу"])
    ws.append(["акъыл + эт-джен (словарь-1)",
               "в цифровых выгрузках их нет; нужен скан/фото страниц печатного словаря-1"])
    ws.append(["словарь омонимов (sources/3), синонимов (sources/4)",
               "жду файлы — без них пп. 4–5 дорожной карты стоят"])
    style(ws, {"A": 45, "B": 80}, "AB")

    # ---------- 9. Отклонено ----------
    ws = wb.create_sheet("9_Отклонено_ру-кб")
    ws.append(["форма", "вердикт триажа", "базовое слово", "причина скептика", "cite", "ctx"])
    for r in jl(RQ / "rus_triage_rejected.jsonl"):
        ws.append([r["form"], r["verdict"], r.get("base", ""),
                   r.get("correction_reason", ""), r["count_cite"], r["count_ctx"]])
    style(ws, {"A": 24, "B": 18, "C": 20, "D": 60, "E": 6, "F": 6}, "D")

    out = RQ / "manual_review_2.xlsx"
    wb.save(out)
    print(f"OK -> {out}")
    print(f"needs_review новых: {n_new}; ◊-примеров: {min(30, len(dia))} из {len(dia)}; "
          f"пометы-хвост: {len(tags)}")


if __name__ == "__main__":
    main()
